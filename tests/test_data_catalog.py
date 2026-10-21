"""Unit tests for Hyperflix Data Discovery and Data Catalog Generator."""

from __future__ import annotations

import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd
from bson import ObjectId

from scripts import generar_catalogo


class DataCatalogTests(unittest.TestCase):
    def test_architectural_classification(self) -> None:
        # OLTP (Transaccional)
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("usuarios"),
            "OLTP (Transaccional)",
        )
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("peliculas"),
            "OLTP (Transaccional)",
        )
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("suscripciones"),
            "OLTP (Transaccional)",
        )
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("pagos"),
            "OLTP (Transaccional)",
        )
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("reproducciones"),
            "OLTP (Transaccional)",
        )

        # OLAP (Analítico)
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("reproducciones_analitica"),
            "OLAP (Analítico)",
        )
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("consumo_analitico"),
            "OLAP (Analítico)",
        )

        # Data Warehouse (Modelo Estrella)
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("dim_tiempo"),
            "Data Warehouse (Modelo Estrella)",
        )
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("dim_usuario"),
            "Data Warehouse (Modelo Estrella)",
        )
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("dim_contenido"),
            "Data Warehouse (Modelo Estrella)",
        )
        self.assertEqual(
            generar_catalogo.classify_collection_architecture("fact_visualizaciones"),
            "Data Warehouse (Modelo Estrella)",
        )

    def test_recursive_json_flattening(self) -> None:
        sample_doc = {
            "_id": ObjectId("507f1f77bcf86cd799439011"),
            "usuario": {
                "id": "USR001",
                "perfil": {
                    "nivel": "VIP",
                    "idioma": "es",
                },
            },
            "dispositivo": {
                "tipo": "Smart TV",
                "os": "Android",
            },
            "calidad": "4K",
        }

        flattened = generar_catalogo.flatten_document(sample_doc)

        self.assertIn("_id", flattened)
        self.assertIn("usuario.id", flattened)
        self.assertIn("usuario.perfil.nivel", flattened)
        self.assertIn("usuario.perfil.idioma", flattened)
        self.assertIn("dispositivo.tipo", flattened)
        self.assertIn("dispositivo.os", flattened)
        self.assertIn("calidad", flattened)
        self.assertEqual(flattened["usuario.perfil.nivel"], "VIP")
        self.assertEqual(flattened["dispositivo.os"], "Android")

    def test_type_translation(self) -> None:
        self.assertEqual(generar_catalogo.translate_bson_type(ObjectId()), "String (UUID)")
        self.assertEqual(generar_catalogo.translate_bson_type(datetime.now(timezone.utc)), "Timestamp")
        self.assertEqual(generar_catalogo.translate_bson_type(True), "Boolean")
        self.assertEqual(generar_catalogo.translate_bson_type(120), "Integer")
        self.assertEqual(generar_catalogo.translate_bson_type(29.99), "Double")
        self.assertEqual(generar_catalogo.translate_bson_type(["Drama", "4K"]), "Array")
        self.assertEqual(generar_catalogo.translate_bson_type("Hyperflix"), "String")
        self.assertEqual(generar_catalogo.translate_bson_type(None), "Null / Desconocido")

    def test_key_and_index_detection(self) -> None:
        # Primary Key
        self.assertEqual(generar_catalogo.is_primary_key("_id"), "Sí")
        self.assertEqual(generar_catalogo.is_primary_key("usuario_id"), "No")

        # Foreign Key
        self.assertEqual(generar_catalogo.is_foreign_key("usuario_id"), "Sí")
        self.assertEqual(generar_catalogo.is_foreign_key("contenido_id"), "Sí")
        self.assertEqual(generar_catalogo.is_foreign_key("suscripcion.id"), "Sí")
        self.assertEqual(generar_catalogo.is_foreign_key("_id"), "No")
        self.assertEqual(generar_catalogo.is_foreign_key("duracion_minutos"), "No")

        # Indexes
        index_keys = {"codigo_usuario", "email", "fecha_id.anio"}
        self.assertEqual(generar_catalogo.is_indexed_field("codigo_usuario", index_keys), "Sí")
        self.assertEqual(generar_catalogo.is_indexed_field("fecha_id.anio", index_keys), "Sí")
        self.assertEqual(generar_catalogo.is_indexed_field("duracion_minutos", index_keys), "No")

    def test_semantic_business_descriptions(self) -> None:
        desc_user = generar_catalogo.get_field_business_description("usuario_id")
        self.assertIn("suscriptor", desc_user.lower())

        desc_minutes = generar_catalogo.get_field_business_description("minutos_reproducidos")
        self.assertIn("minutos", desc_minutes.lower())

        desc_device = generar_catalogo.get_field_business_description("dispositivo.os")
        self.assertIn("sistema operativo", desc_device.lower())

    def test_excel_export_creates_valid_workbook_and_tabs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_file = Path(tmp_dir) / "Test_Catalogo_Hyperflix.xlsx"

            df_resumen = pd.DataFrame([
                {
                    "Colección": "usuarios",
                    "Zona Arquitectónica (OLTP/OLAP/DW)": "OLTP (Transaccional)",
                    "Total Registros (Docs)": 5000,
                    "Tamaño Estimado": "1.25 MB",
                },
                {
                    "Colección": "fact_visualizaciones",
                    "Zona Arquitectónica (OLTP/OLAP/DW)": "Data Warehouse (Modelo Estrella)",
                    "Total Registros (Docs)": 24005,
                    "Tamaño Estimado": "4.80 MB",
                },
            ])

            df_diccionario = pd.DataFrame([
                {
                    "Colección": "usuarios",
                    "Campo": "email",
                    "Tipo de Dato": "String",
                    "Es Primary Key (Sí/No)": "No",
                    "Es Foreign Key (Sí/No)": "No",
                    "Está Indexado (Sí/No)": "Sí",
                    "Descripción de Negocio": "Correo electrónico del suscriptor.",
                }
            ])

            result_path = generar_catalogo.export_catalog_to_excel(
                df_resumen, df_diccionario, output_file
            )

            self.assertTrue(result_path.exists())

            # Load and verify with openpyxl
            wb = openpyxl.load_workbook(result_path)
            self.assertIn("Resumen Arquitectónico", wb.sheetnames)
            self.assertIn("Diccionario de Datos", wb.sheetnames)

            ws_resumen = wb["Resumen Arquitectónico"]
            self.assertEqual(ws_resumen.cell(row=1, column=1).value, "Colección")
            self.assertEqual(ws_resumen.cell(row=2, column=1).value, "usuarios")

            ws_dict = wb["Diccionario de Datos"]
            self.assertEqual(ws_dict.cell(row=1, column=2).value, "Campo")
            self.assertEqual(ws_dict.cell(row=2, column=2).value, "email")


if __name__ == "__main__":
    unittest.main()
